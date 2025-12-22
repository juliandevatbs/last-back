import json
import logging
import os
import traceback

import win32com.client as win32
from docx import Document
from openpyxl.reader.excel import load_workbook

from core.exceptions import NotJsonConfigFound
from graphics.Services.copy.copy_graphics import CopyGraphics
from graphics.Services.read.read_graphics_data import ReadGraphicsData
from graphics.Services.write.write_insitu_data import WriteInsituData
from graphics.Services.write.write_parameters_results import WriteParametersResults

logger = logging.getLogger(__name__)

class GraphicGenerator:

    def __init__(self):


        self.graphics_config = None
        self.chain_of_custody = None


        #Constants
        self.JSON_FOLDER_URL = "fields_config/"
        self.word_path = "templates/informe_final.docx"

        # Excels
        self.graphics_obj = None
        self.chain_custody_obj = None
        self.json_config = None
        self.word_obj = None

        # COM objects
        self.graphics_excel_com = None

        #Instances
        self.reader_graphics_data = ReadGraphicsData()
        self.writer_graphics_data = WriteInsituData()
        self.writer_parameters_data = WriteParametersResults()
        self.copy_graphics = None

    def search_graphics_config(self, report_name: str):

        report_name = "GRAFICAS_" + report_name + ".json"

        if report_name not in os.listdir(self.JSON_FOLDER_URL):
            raise NotJsonConfigFound(f"The configuration file for the {report_name} graphics does not exist")

        config_path = os.path.join(self.JSON_FOLDER_URL, report_name)

        try:

            with open(config_path, 'r', encoding='utf-8') as file:

                self.graphics_config = json.load(file)

            logger.info(f"Grapchis config loaded succesfully {report_name}")

        except json.JSONDecodeError as e:

            logger.error(f"Invalid json in {report_name}: {e}")
            raise

        except Exception as e:

            logger.error(f"Error reading {report_name}  graphics config: {e}")
            raise

    def load_graphics_excel(self, url_graphics_config, read_only):

        if not url_graphics_config:

            logger.error(f"No excel url for the charts is provided")

        graphics_excel = load_workbook(url_graphics_config, data_only=read_only)

        self.graphics_obj = graphics_excel

        logger.info("Graphics excel open succesfully")

    def load_graphics_excel_com(self, excel_path):

        try:
            excel = win32.Dispatch("Excel.Application")
            excel.Visible = False
            wb = excel.Workbooks.Open(os.path.abspath(excel_path))
            self.graphics_excel_com = wb
            logger.info("Excel COM loaded successfully for graphics")
        except Exception as e:
            logger.error(f"Error loading Excel COM: {e}", exc_info=True)
            raise


    def load_json_config(self, report_name:str):


        path_config = f"{self.JSON_FOLDER_URL}GRAFICAS_{report_name}.json"

        try:
            
            with open(path_config, 'r', encoding='utf-8') as archivo:

                self.json_config = json.load(archivo)

        except Exception as ex:

            logger.error(f"Error opening json {ex}")

        return

    def load_chain_of_custody(self, chain_custody):


        if not chain_custody:

            logger.error("Not chain of custody excel")

            return

        try:

            self.chain_of_custody = chain_custody

        except Exception as ex:

            logger.error(f"Error opening the chain of custody excel")

    def load_word_obj(self):

        if not self.word_path:

            logger.error("No se suministro ruta del word")

            raise ValueError("Word path is None")
        try:

            self.word_obj = Document(self.word_path)

        except Exception as ex:

            logger.error(f"Error cargando documento {ex}", exc_info=True)

            raise

    def copy_graphics_proc(self):
        try:
            if self.copy_graphics is None:


                self.copy_graphics = CopyGraphics(
                    self.graphics_excel_com,
                    self.json_config,
                    self.word_obj,
                    self.word_path
                )

            self.copy_graphics.copy_and_insert_editable_charts()
            self.copy_graphics.paste_graphics()

        except Exception as ex:
            logger.error(f"Error copiando graficas", exc_info=True)
            raise

    def load_word_obj_com(self):
        """Abre Word con COM para poder pegar gráficos"""
        if not self.word_path:
            logger.error("No se suministro ruta del word")
            raise ValueError("Word path is None")

        try:
            word = win32.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(os.path.abspath(self.word_path))
            self.word_obj = doc
            logger.info("Word COM loaded successfully")
        except Exception as ex:
            logger.error(f"Error cargando documento Word con COM: {ex}", exc_info=True)
            raise


    def load_objs(self, report_name, chain_custody, read_only: False):
        try:
            self.load_json_config(report_name)
            self.load_graphics_excel(self.json_config["RUTA_EXCEL_GRAFICAS"], read_only)

            if self.graphics_obj and not read_only:
                self.graphics_obj.close()
                logger.info("Closed openpyxl workbook before COM access")

            self.load_graphics_excel_com(self.json_config["RUTA_EXCEL_GRAFICAS"])
            self.load_chain_of_custody(chain_custody)
            self.load_word_obj()

        except Exception as ex:
            logger.error(f"Error in the main thread {ex}")
            logger.error(f"Traceback: {traceback.format_exc()}")

    def main(self, report_name: str, chain_custody):
        try:

            self.load_json_config(report_name)
            self.load_graphics_excel(self.json_config["RUTA_EXCEL_GRAFICAS"], False)
            self.load_chain_of_custody(chain_custody)

            self.reader_graphics_data.get_chain_custody(chain_custody)
            self.reader_graphics_data.load_json_config(self.json_config)
            self.reader_graphics_data.get_cod_muestras()
            sampler_parameters, parameters_results = self.reader_graphics_data.read_parameters_to_sampler(
                self.graphics_obj)

            tables_data = self.reader_graphics_data.read_table_ph_solidos_caudales()
            formatted_tables_data = self.reader_graphics_data.complete_table_data()
            self.writer_graphics_data.load_insitu_data(formatted_tables_data)
            self.writer_graphics_data.load_json_config(self.json_config)
            self.writer_graphics_data.graphics_excel = self.graphics_obj
            self.writer_graphics_data.write_insitu_data()
            self.writer_parameters_data.write_parameters_result(self.graphics_obj, self.json_config, parameters_results)

            self.graphics_obj.save(self.json_config["RUTA_EXCEL_GRAFICAS"])
            self.graphics_obj.close()
            self.graphics_obj = None
            logger.info("Phase 1: Excel saved and closed")

            self.load_graphics_excel(self.json_config["RUTA_EXCEL_GRAFICAS"], True)
            final_results = self.reader_graphics_data.graphics_post_calculus(self.graphics_obj, self.json_config)
            self.graphics_obj.close()
            self.graphics_obj = None
            logger.info("Phase 2: Calculations completed and closed")

            self.load_graphics_excel(self.json_config["RUTA_EXCEL_GRAFICAS"], False)
            self.writer_parameters_data.write_final_results(final_results, self.graphics_obj, self.json_config)

            self.graphics_obj.save(self.json_config["RUTA_EXCEL_GRAFICAS"])
            self.graphics_obj.close()
            self.graphics_obj = None
            logger.info("Phase 3: Final results saved and closed")

            import time
            time.sleep(1)

            self.load_word_obj_com()
            self.load_graphics_excel_com(self.json_config["RUTA_EXCEL_GRAFICAS"])
            self.copy_graphics_proc()

            if self.graphics_excel_com:
                self.graphics_excel_com.Close(SaveChanges=False)
                excel_app = self.graphics_excel_com.Application
                excel_app.Quit()
                logger.info("COM Excel closed")
                self.graphics_excel_com = None

            self.word_obj.save(self.word_path)
            logger.info("Word document saved - Process completed")

        except Exception as ex:
            logger.error(f"Error in the main thread {ex}")
            logger.error(f"Traceback: {traceback.format_exc()}")

            try:
                if self.graphics_obj:
                    self.graphics_obj.close()
                if self.graphics_excel_com:
                    self.graphics_excel_com.Close(SaveChanges=False)
                    self.graphics_excel_com.Application.Quit()
            except:
                pass










